from __future__ import annotations

import mimetypes
import random
import re
import shutil
import tempfile
import time
import uuid
from functools import wraps
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def create_excel(data: list[dict], output_path: str | Path) -> str:
    """Write a styled Excel report, preserving the original behaviour."""
    df = pd.DataFrame(data)
    MAX_WIDTH = 70

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        sheet_name = "Sheet1"
        df.to_excel(writer, index=False, sheet_name=sheet_name)

        worksheet = writer.sheets[sheet_name]

        header_fill = PatternFill(start_color="185C37", end_color="185C37", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        for i, column in enumerate(df.columns):
            # pandas >= 3 keeps missing values as NA when casting to string and
            # ``len()`` raises on those; ``str.len()`` returns NA instead.
            max_data_len = df[column].astype("string").str.len().max()
            header_len = len(str(column))

            if pd.isna(max_data_len):
                max_data_len = 0

            ideal_width = max(max_data_len, header_len) + 2
            adjusted_width = min(ideal_width, MAX_WIDTH)

            col_letter = get_column_letter(i + 1)
            worksheet.column_dimensions[col_letter].width = adjusted_width

    return str(output_path)


def get_mime_type_from_file(
    file_path: str | Path, default_mime_type: str = "application/pdf"
) -> str:
    path = Path(file_path)
    mime_type, _ = mimetypes.guess_type(path)

    safe_mapping = {
        ".pdf": "application/pdf",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
    }

    if mime_type not in ["application/pdf", "image/jpeg", "image/png"]:
        return safe_mapping.get(path.suffix.lower(), default_mime_type)

    return mime_type


def jitter_call(low: float = 1, high: float = 3):
    """Small randomized delay used to avoid hammering the LLM API."""

    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            time.sleep(random.uniform(low, high))
            return func(*args, **kwargs)

        return wrapper

    return decorator


def extract_json_from_llm_response(response_text: str | None) -> str | None:
    if not response_text:
        return None

    cleaned_text = response_text.strip()
    match = re.search(r"```(?:json)?\s*({.*?})\s*```", cleaned_text, re.DOTALL)

    if match:
        return match.group(1)

    try:
        start_index = cleaned_text.index("{")
        end_index = cleaned_text.rindex("}")
        if start_index < end_index:
            return cleaned_text[start_index : end_index + 1]
    except ValueError:
        return None
    return None


def make_workdir(prefix: str = "pfasimplu_") -> Path:
    return Path(tempfile.mkdtemp(prefix=prefix))


def safe_unlink(path: str | Path | None) -> None:
    if not path:
        return
    try:
        Path(path).unlink(missing_ok=True)
    except OSError:
        pass


def safe_rmtree(path: str | Path | None) -> None:
    if not path:
        return
    shutil.rmtree(path, ignore_errors=True)


def unique_name(prefix: str, suffix: str) -> str:
    return f"{prefix}{uuid.uuid4().hex}{suffix}"
